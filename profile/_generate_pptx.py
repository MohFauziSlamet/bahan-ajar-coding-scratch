#!/usr/bin/env python3
"""
Generator PPTX perkenalan pengajar + sesi perkenalan siswa.

Sumber isi : moh_fauzi_slamet_cv.pdf + foto venturo-k2-*.jpg di folder ini.
Tema/layout: dipinjam dari ../tab_code/_generate_pptx.py agar seragam dengan
             dek bahan ajar Scratch lainnya.

Jalankan:  python3 _generate_pptx.py
Butuh   :  pip install python-pptx
"""

import importlib.util
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT_DIR = Path(__file__).resolve().parent
KIT_PATH = OUT_DIR.parent / "tab_code" / "_generate_pptx.py"
if not KIT_PATH.exists():
    sys.exit(f"Tidak menemukan kit layout di {KIT_PATH}")


def _load_kit():
    """Muat kit layout tab_code dengan nama modul berbeda (hindari tabrakan nama)."""
    spec = importlib.util.spec_from_file_location("scratch_pptx_kit", KIT_PATH)
    mod = importlib.util.module_from_spec(spec)
    sys.modules["scratch_pptx_kit"] = mod
    spec.loader.exec_module(mod)
    return mod


kit = _load_kit()

W, H, MARGIN, BODY_W = kit.W, kit.H, kit.MARGIN, kit.BODY_W
INK, MUTED, WHITE = kit.INK, kit.MUTED, kit.WHITE
CODE_BG, NOTE_BG = kit.CODE_BG, kit.NOTE_BG
FONT, MONO = kit.FONT, kit.MONO
rgb, tint = kit.rgb, kit.tint
_txbox, _para, _rect, _blank = kit._txbox, kit._para, kit._rect, kit._blank
_header, _footer = kit._header, kit._footer

BLUE = rgb("0E7EC4")        # biru Venturo
GREEN = rgb("5FB63A")       # hijau Venturo
ORANGE = rgb("FF8C1A")      # oranye Scratch
PURPLE = rgb("7C4DFF")

DECK = "Perkenalan  ·  Coding Scratch"
IMG_AR = 5184 / 3456        # semua foto profil 3:2
PHOTO = {n: OUT_DIR / f"venturo-k2-{n}.jpg" for n in (193, 194, 195, 196, 197)}


# ------------------------------------------------------------- helper

def _photo(s, path, x, y, w, h, focus=0.5, ar=IMG_AR):
    """Tempel foto memenuhi kotak (cover), kelebihannya dipotong lewat crop."""
    pic = s.shapes.add_picture(str(path), x, y, w, h)
    box_ar = w / h
    if box_ar < ar:                       # foto terlalu lebar → potong kiri/kanan
        cut = 1 - box_ar / ar
        pic.crop_left = cut * focus
        pic.crop_right = cut * (1 - focus)
    else:                                 # foto terlalu tinggi → potong atas/bawah
        cut = 1 - ar / box_ar
        pic.crop_top = cut * focus
        pic.crop_bottom = cut * (1 - focus)
    return pic


def _chip(s, x, y, text, color, w=None, size=13):
    """Label kecil berlatar warna muda."""
    w = w or Inches(2.0)
    sh = _rect(s, x, y, w, Inches(0.42), tint(color, 0.86))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.14)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _para(tf, text, size, bold=True, color=color, first=True, space_after=0,
          align=PP_ALIGN.CENTER)
    return sh


def _card(s, x, y, w, h, color, title, sub, lines, metric=None):
    """Kartu isi: judul, sub-judul, beberapa baris, plus angka penting."""
    _rect(s, x, y, w, h, tint(color, 0.94))
    _rect(s, x, y, Inches(0.09), h, color, shape=MSO_SHAPE.RECTANGLE)

    tf = _txbox(s, x + Inches(0.34), y + Inches(0.26), w - Inches(0.68), h - Inches(0.5))
    _para(tf, title, 19, bold=True, color=INK, first=True, space_after=2)
    _para(tf, sub, 12, color=MUTED, space_after=8)
    if metric:
        _para(tf, metric, 14, bold=True, color=color, space_after=8)
    for line in lines:
        _para(tf, line, 13, color=INK, space_after=4)


# -------------------------------------------------------- jenis slide

def slide_cover(prs):
    s = _blank(prs)
    split = Inches(6.9)

    panel = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, split, H)
    panel.fill.solid()
    panel.fill.fore_color.rgb = BLUE
    panel.line.fill.background()
    panel.shadow.inherit = False

    _photo(s, PHOTO[193], split, 0, W - split, H, focus=0.42)

    tk = _txbox(s, MARGIN, Inches(1.55), split - MARGIN - Inches(0.6), Inches(0.3))
    _para(tk, "BAHAN AJAR CODING SCRATCH", 13, bold=True, color=tint(BLUE, 0.75),
          first=True, space_after=0, line_spacing=1.0)

    tt = _txbox(s, MARGIN, Inches(2.15), split - MARGIN - Inches(0.6), Inches(1.6))
    _para(tt, "Perkenalan", 44, bold=True, color=WHITE, first=True,
          space_after=0, line_spacing=1.0)

    tn = _txbox(s, MARGIN, Inches(3.15), split - MARGIN - Inches(0.6), Inches(1.2))
    _para(tn, "Moh Fauzi Slamet", 30, bold=True, color=WHITE, first=True,
          space_after=4, line_spacing=1.05)
    _para(tn, "Flutter Mobile Developer", 18, color=tint(BLUE, 0.82),
          space_after=0, line_spacing=1.05)

    tb = _txbox(s, MARGIN, Inches(5.55), split - MARGIN - Inches(0.6), Inches(0.9))
    _para(tb, "PT Venturo Pro Indonesia  ·  Malang", 14, color=tint(BLUE, 0.78),
          first=True, space_after=0, line_spacing=1.1)
    return s


def slide_halo(prs, page):
    s = _blank(prs)
    _header(s, BLUE, "Halo, kenalkan saya Kak Fauzi", "SIAPA YANG MENGAJAR KALIAN")

    _photo(s, PHOTO[195], MARGIN, Inches(1.55), Inches(5.1), Inches(4.25), focus=0.45)

    x = MARGIN + Inches(5.5)
    w = BODY_W - Inches(5.5)
    tf = _txbox(s, x, Inches(1.62), w, Inches(4.2))
    _para(tf, "TENTANG SAYA", 11, bold=True, color=MUTED, first=True, space_after=8)
    for line in [
        "Nama lengkap: Moh Fauzi Slamet — panggil saja Kak Fauzi.",
        "Tinggal di Malang, Jawa Timur.",
        "Lulusan S.Kom Institut Teknologi dan Bisnis ASIA Malang (2018–2022).",
        "Bekerja sebagai Mobile Developer di PT Venturo Pro Indonesia sejak Agustus 2022.",
        "Hampir 4 tahun membuat aplikasi HP untuk klien pemerintah, kesehatan, donasi, dan komunitas.",
    ]:
        _para(tf, "•  " + line, 15, color=INK, space_after=10)

    ny = Inches(5.95)
    _rect(s, x, ny, w, Inches(0.85), NOTE_BG)
    tfn = _txbox(s, x + Inches(0.22), ny + Inches(0.16), w - Inches(0.44), Inches(0.6))
    _para(tfn, "Pekerjaan saya sehari-hari: menulis kode agar sebuah aplikasi bisa "
               "dipakai orang banyak di HP mereka.", 13, color=INK, first=True, space_after=0)

    _footer(s, DECK, page)
    return s


def slide_pekerjaan(prs, page):
    s = _blank(prs)
    _header(s, GREEN, "Apa itu Mobile Developer?", "PEKERJAAN SAYA")

    tf = _txbox(s, MARGIN, Inches(1.55), BODY_W, Inches(0.7))
    _para(tf, "Orang yang membuat aplikasi di HP — seperti aplikasi yang ada di layar HP kalian sekarang.",
          20, bold=True, color=INK, first=True, space_after=0, line_spacing=1.15)

    items = [
        ("Merancang tampilan", "Tombol, halaman, warna, animasi — supaya enak dilihat dan mudah dipakai."),
        ("Menyambung ke internet", "Mengambil data dari server: login, daftar produk, chat, pembayaran."),
        ("Menjaga aplikasi tetap sehat", "Memperbaiki error agar aplikasi tidak macet walau dipakai jutaan orang."),
    ]
    gap = Inches(0.3)
    cw = (BODY_W - 2 * gap) / 3
    for i, (t, d) in enumerate(items):
        x = MARGIN + i * (cw + gap)
        _rect(s, x, Inches(2.55), cw, Inches(2.35), tint(GREEN, 0.93))
        tfc = _txbox(s, x + Inches(0.3), Inches(2.8), cw - Inches(0.6), Inches(1.9))
        _para(tfc, str(i + 1), 26, bold=True, color=GREEN, first=True, space_after=4)
        _para(tfc, t, 17, bold=True, color=INK, space_after=6)
        _para(tfc, d, 13, color=MUTED, space_after=0, line_spacing=1.15)

    ny = Inches(5.35)
    _rect(s, MARGIN, ny, BODY_W, Inches(1.15), NOTE_BG)
    tfn = _txbox(s, MARGIN + Inches(0.28), ny + Inches(0.2), BODY_W - Inches(0.56), Inches(0.8))
    _para(tfn, "MIRIP SCRATCH, KAN?", 11, bold=True, color=rgb("B26A00"), first=True, space_after=5)
    _para(tfn, "Di Scratch ada Panggung (tampilan) dan blok kode (perintah). Di pekerjaan saya juga sama — "
               "hanya saja perintahnya tidak berbentuk blok warna-warni, tetapi ditulis sebagai teks.",
          14, color=INK, space_after=0)

    _footer(s, DECK, page)
    return s


def slide_alat(prs, page):
    s = _blank(prs)
    _header(s, PURPLE, "Alat dan bahasa yang saya pakai", "ISI TAS KERJA SAYA")

    groups = [
        ("BAHASA & FRAMEWORK", ["Dart", "Flutter"], PURPLE),
        ("PENGATUR DATA APLIKASI", ["Provider", "GetX", "BLoC", "Riverpod"], BLUE),
        ("PENYAMBUNG & PENDUKUNG", ["REST API", "Firebase", "Figma", "Bitbucket Pipelines"], GREEN),
    ]
    y = Inches(1.6)
    for label, chips, color in groups:
        tf = _txbox(s, MARGIN, y, BODY_W, Inches(0.3))
        _para(tf, label, 11, bold=True, color=MUTED, first=True, space_after=0)
        cx = MARGIN
        for c in chips:
            cw = Inches(0.42 + 0.135 * len(c))
            _chip(s, cx, y + Inches(0.38), c, color, w=cw, size=14)
            cx += cw + Inches(0.22)
        y += Inches(1.25)

    tf2 = _txbox(s, MARGIN, Inches(5.35), BODY_W, Inches(0.5))
    _para(tf2, "Semua nama di atas hanyalah ALAT. Yang paling penting bukan alatnya, tapi cara berpikirnya.",
          17, bold=True, color=INK, first=True, space_after=0)

    ny = Inches(6.0)
    _rect(s, MARGIN, ny, BODY_W, Inches(0.85), NOTE_BG)
    tfn = _txbox(s, MARGIN + Inches(0.28), ny + Inches(0.17), BODY_W - Inches(0.56), Inches(0.6))
    _para(tfn, "Cara berpikir itu yang kita latih di Scratch: urutan, perulangan, kondisi, dan kejadian (event).",
          14, color=INK, first=True, space_after=0)

    _footer(s, DECK, page)
    return s


def slide_karya(prs, page):
    s = _blank(prs)
    _header(s, BLUE, "Aplikasi yang pernah saya kerjakan", "4 APLIKASI, SEMUA DIPAKAI ORANG SUNGGUHAN")

    cards = [
        (BLUE, "Pusaka Super Apps", "Kementerian Agama RI",
         "1.000.000+ unduhan di Google Play",
         ["Al-Qur'an digital, jadwal salat, arah kiblat,", "kalender hijriah, daftar nikah, sertifikasi halal."]),
        (GREEN, "Hayyu Apps", "Hayyu Skin Clinic  ·  Google Play & App Store",
         "Rating 4,9 dari 5 di App Store",
         ["Konsultasi dokter lewat chat & video call,", "reservasi perawatan, riwayat pengobatan."]),
        (ORANGE, "Lazisnu Apps", "NU Care–LAZISNU",
         "Zakat, infak, sedekah, fidyah, kurban",
         ["Alur donasi bertahap dan pembayaran lewat", "Virtual Account. Kini tidak lagi ada di store."]),
        (PURPLE, "CUiT: Community Enabler", "Produk komunitas  ·  Google Play & App Store",
         "Rating 4,8 dari 5 di App Store",
         ["Misi/tugas komunitas, konten interaktif,", "chat waktu-nyata, dan hadiah."]),
    ]
    gap = Inches(0.32)
    cw = (BODY_W - gap) / 2
    ch = Inches(2.25)
    for i, (color, title, sub, metric, lines) in enumerate(cards):
        x = MARGIN + (i % 2) * (cw + gap)
        y = Inches(1.5) + (i // 2) * (ch + Inches(0.28))
        _card(s, x, y, cw, ch, color, title, sub, lines, metric=metric)

    tf = _txbox(s, MARGIN, Inches(6.55), BODY_W, Inches(0.3))
    _para(tf, "Angka store per Juli 2026. Semua dikerjakan bersama tim di PT Venturo Pro Indonesia.",
          11, color=MUTED, first=True, space_after=0)

    _footer(s, DECK, page)
    return s


def slide_pusaka(prs, page):
    s = _blank(prs)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    band.fill.solid()
    band.fill.fore_color.rgb = BLUE
    band.line.fill.background()
    band.shadow.inherit = False

    tk = _txbox(s, MARGIN, Inches(1.35), BODY_W, Inches(0.35))
    _para(tk, "APLIKASI YANG PALING BANYAK DIPAKAI", 14, bold=True,
          color=tint(BLUE, 0.75), first=True, space_after=0, line_spacing=1.0)

    tn = _txbox(s, MARGIN, Inches(2.0), BODY_W, Inches(1.5))
    _para(tn, "1.000.000+", 72, bold=True, color=WHITE, first=True,
          space_after=0, line_spacing=1.0)

    tt = _txbox(s, MARGIN, Inches(3.5), BODY_W, Inches(1.6))
    _para(tt, "unduhan aplikasi Pusaka di Google Play — aplikasi resmi Kementerian Agama RI.",
          26, color=WHITE, first=True, space_after=0, line_spacing=1.25)

    tb = _txbox(s, MARGIN, Inches(5.4), BODY_W, Inches(1.2))
    _para(tb, "Artinya: kode yang saya tulis bareng tim dibuka oleh sejuta orang lebih di HP mereka. "
              "Satu tombol yang salah pun akan langsung terasa oleh banyak orang.",
          17, color=tint(BLUE, 0.85), first=True, space_after=0, line_spacing=1.25)
    return s


def slide_perjalanan(prs, page):
    s = _blank(prs)
    _header(s, ORANGE, "Perjalanan saya sampai jadi developer", "TIDAK ADA YANG INSTAN")

    steps = [
        ("2018", "Masuk kuliah", "Institut Teknologi dan Bisnis ASIA Malang."),
        ("2022", "Lulus S.Kom", "Selesai kuliah setelah 4 tahun belajar."),
        ("Agu 2022", "Kerja pertama", "Diterima sebagai Mobile Developer di PT Venturo Pro Indonesia."),
        ("2026", "Hampir 4 tahun", "4 aplikasi produksi, 4 cara mengatur data aplikasi dikuasai."),
    ]
    gap = Inches(0.28)
    cw = (BODY_W - 3 * gap) / 4
    y = Inches(1.9)
    _rect(s, MARGIN, y + Inches(0.62), BODY_W, Inches(0.06), tint(ORANGE, 0.65),
          shape=MSO_SHAPE.RECTANGLE)
    for i, (year, title, desc) in enumerate(steps):
        x = MARGIN + i * (cw + gap)
        _chip(s, x, y + Inches(0.42), year, ORANGE, w=Inches(1.7), size=14)
        tf = _txbox(s, x, y + Inches(1.25), cw, Inches(2.2))
        _para(tf, title, 18, bold=True, color=INK, first=True, space_after=6)
        _para(tf, desc, 13, color=MUTED, space_after=0, line_spacing=1.2)

    ny = Inches(5.5)
    _rect(s, MARGIN, ny, BODY_W, Inches(1.15), tint(ORANGE, 0.92))
    tfn = _txbox(s, MARGIN + Inches(0.3), ny + Inches(0.2), BODY_W - Inches(0.6), Inches(0.8))
    _para(tfn, "Pesan untuk kalian", 13, bold=True, color=ORANGE, first=True, space_after=5)
    _para(tfn, "Saya juga mulai dari nol dan sering salah. Yang membedakan bukan bakat, "
               "tapi mau mencoba lagi setelah kodenya error.", 16, color=INK, space_after=0)

    _footer(s, DECK, page)
    return s


def slide_jembatan(prs, page):
    s = _blank(prs)
    _header(s, GREEN, "Dari blok Scratch ke kode sungguhan", "APA GUNANYA BELAJAR SCRATCH?")

    rows = [
        ("when green flag clicked", "Event — kode berjalan saat aplikasi dibuka."),
        ("repeat 10", "Perulangan — menampilkan daftar 10 produk."),
        ("if ... then ... else", "Percabangan — password benar atau salah."),
        ("set [skor] to 0", "Variabel — menyimpan data pengguna."),
        ("broadcast [pesan]", "Kirim pesan antar bagian aplikasi."),
        ("sprite", "Widget — potongan tampilan di layar."),
    ]
    y = Inches(1.55)
    rh = Inches(0.72)
    lw = Inches(5.0)
    for i, (blok, arti) in enumerate(rows):
        if i % 2 == 0:
            _rect(s, MARGIN, y, BODY_W, rh, kit.ROW_ALT, shape=MSO_SHAPE.RECTANGLE)
        tfl = _txbox(s, MARGIN + Inches(0.25), y + Inches(0.2), lw, Inches(0.4))
        _para(tfl, blok, 15, color=INK, font=MONO, first=True, space_after=0)
        tfr = _txbox(s, MARGIN + lw + Inches(0.6), y + Inches(0.2), BODY_W - lw - Inches(0.85), Inches(0.4))
        _para(tfr, arti, 15, color=INK, first=True, space_after=0)
        y += rh

    ny = Inches(6.0)
    _rect(s, MARGIN, ny, BODY_W, Inches(0.85), tint(GREEN, 0.9))
    tfn = _txbox(s, MARGIN + Inches(0.28), ny + Inches(0.17), BODY_W - Inches(0.56), Inches(0.6))
    _para(tfn, "Konsepnya persis sama. Kalau kalian paham blok-blok ini, kalian sudah setengah jalan "
               "menuju bahasa pemrograman sungguhan.", 15, bold=True, color=INK, first=True, space_after=0)

    _footer(s, DECK, page)
    return s


def slide_giliran(prs):
    s = _blank(prs)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    band.fill.solid()
    band.fill.fore_color.rgb = ORANGE
    band.line.fill.background()
    band.shadow.inherit = False

    tk = _txbox(s, MARGIN, Inches(0.75), BODY_W, Inches(0.35))
    _para(tk, "SESI PERKENALAN", 14, bold=True, color=tint(ORANGE, 0.8),
          first=True, space_after=0, line_spacing=1.0)

    tt = _txbox(s, MARGIN, Inches(1.25), BODY_W, Inches(0.9))
    _para(tt, "Sekarang giliran kalian!", 44, bold=True, color=WHITE, first=True,
          space_after=0, line_spacing=1.0)

    qs = [
        ("1", "Siapa nama kalian?", "Nama lengkap dan nama panggilan."),
        ("2", "Dari kelas apa?", "Sebutkan kelas kalian sekarang."),
        ("3", "Mengapa ingin belajar coding?", "Alasan kalian mau ada di kelas ini."),
        ("4", "Motivasi tertinggi kalian apa?", "Hal yang membuat kalian tetap belajar walau susah."),
    ]
    y = Inches(2.6)
    rh = Inches(1.0)
    for num, q, hint in qs:
        _rect(s, MARGIN, y, BODY_W, Inches(0.88), WHITE)
        tfn = _txbox(s, MARGIN + Inches(0.35), y + Inches(0.2), Inches(0.6), Inches(0.5))
        _para(tfn, num, 24, bold=True, color=ORANGE, first=True, space_after=0)
        tfq = _txbox(s, MARGIN + Inches(1.05), y + Inches(0.14), Inches(6.0), Inches(0.6))
        _para(tfq, q, 21, bold=True, color=INK, first=True, space_after=0)
        tfh = _txbox(s, MARGIN + Inches(7.2), y + Inches(0.22), BODY_W - Inches(7.6), Inches(0.5))
        _para(tfh, hint, 13, color=MUTED, first=True, space_after=0)
        y += rh
    return s


def slide_cara_jawab(prs, page):
    s = _blank(prs)
    _header(s, PURPLE, "Cara menjawabnya", "SUPAYA CEPAT DAN SERU")

    tf = _txbox(s, MARGIN, Inches(1.6), Inches(6.6), Inches(0.4))
    _para(tf, "CONTOH JAWABAN", 11, bold=True, color=MUTED, first=True, space_after=8)

    _rect(s, MARGIN, Inches(2.05), Inches(6.6), Inches(3.55), CODE_BG)
    tfc = _txbox(s, MARGIN + Inches(0.3), Inches(2.3), Inches(6.0), Inches(3.1))
    for i, line in enumerate([
        "“Halo, nama saya Rani, panggil saja Rani.",
        "",
        "Saya dari kelas 8B.",
        "",
        "Saya ingin belajar coding karena saya suka main game",
        "dan penasaran bagaimana cara membuatnya.",
        "",
        "Motivasi tertinggi saya: suatu hari saya ingin membuat",
        "game sendiri lalu dimainkan teman-teman saya.”",
    ]):
        _para(tfc, line if line else " ", 15, color=INK, first=(i == 0),
              space_after=2, line_spacing=1.15)

    x = MARGIN + Inches(7.1)
    w = BODY_W - Inches(7.1)
    tfr = _txbox(s, x, Inches(1.6), w, Inches(0.4))
    _para(tfr, "ATURAN MAIN", 11, bold=True, color=MUTED, first=True, space_after=8)

    rules = [
        ("Berdiri", "Supaya semua teman bisa melihat."),
        ("Cukup 1 menit", "Singkat saja, tidak perlu panjang."),
        ("Suara jelas", "Yang di belakang juga harus dengar."),
        ("Jujur", "Tidak ada jawaban yang salah."),
    ]
    y = Inches(2.05)
    for t, d in rules:
        _rect(s, x, y, w, Inches(0.82), tint(PURPLE, 0.93))
        tfi = _txbox(s, x + Inches(0.28), y + Inches(0.13), w - Inches(0.56), Inches(0.6))
        _para(tfi, t, 15, bold=True, color=PURPLE, first=True, space_after=2)
        _para(tfi, d, 12, color=MUTED, space_after=0)
        y += Inches(0.91)

    ny = Inches(5.9)
    _rect(s, MARGIN, ny, Inches(6.6), Inches(0.9), NOTE_BG)
    tfn = _txbox(s, MARGIN + Inches(0.28), ny + Inches(0.18), Inches(6.0), Inches(0.6))
    _para(tfn, "Kak Fauzi akan mencatat motivasi kalian — nanti kita buka lagi di akhir kelas.",
          14, color=INK, first=True, space_after=0)

    _footer(s, DECK, page)
    return s


def slide_penutup(prs):
    s = _blank(prs)
    split = Inches(6.9)

    _photo(s, PHOTO[197], 0, 0, split, H, focus=0.5)

    panel = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, split, 0, W - split, H)
    panel.fill.solid()
    panel.fill.fore_color.rgb = GREEN
    panel.line.fill.background()
    panel.shadow.inherit = False

    x = split + Inches(0.6)
    w = W - split - Inches(1.2)

    tk = _txbox(s, x, Inches(1.6), w, Inches(0.3))
    _para(tk, "SAMPAI JUMPA DI SLIDE BERIKUTNYA", 12, bold=True,
          color=tint(GREEN, 0.8), first=True, space_after=0, line_spacing=1.0)

    tt = _txbox(s, x, Inches(2.15), w, Inches(2.0))
    _para(tt, "Selamat datang di kelas Coding Scratch!", 34, bold=True, color=WHITE,
          first=True, space_after=0, line_spacing=1.1)

    tb = _txbox(s, x, Inches(4.35), w, Inches(1.0))
    _para(tb, "Hari ini kita mulai dari blok. Suatu hari nanti, mungkin dari kalian "
              "lahir aplikasi yang dipakai sejuta orang.", 16, color=tint(GREEN, 0.88),
          first=True, space_after=0, line_spacing=1.25)

    tc = _txbox(s, x, Inches(5.85), w, Inches(1.0))
    _para(tc, "Moh Fauzi Slamet", 15, bold=True, color=WHITE, first=True, space_after=4)
    _para(tc, "ziiidev.vercel.app", 13, color=tint(GREEN, 0.85), space_after=2)
    _para(tc, "github.com/MohFauziSlamet", 13, color=tint(GREEN, 0.85), space_after=0)
    return s


# ---------------------------------------------------------------- main

def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    slide_cover(prs)
    page = 2
    for fn in (slide_halo, slide_pekerjaan, slide_alat, slide_karya):
        fn(prs, page)
        page += 1
    slide_pusaka(prs, page); page += 1
    slide_perjalanan(prs, page); page += 1
    slide_jembatan(prs, page); page += 1
    slide_giliran(prs); page += 1
    slide_cara_jawab(prs, page)
    slide_penutup(prs)

    out = OUT_DIR / "Perkenalan-Pengajar-dan-Siswa.pptx"
    prs.save(out)
    print(f"OK  {out.name}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slide)")


if __name__ == "__main__":
    missing = [p.name for p in PHOTO.values() if not p.exists()]
    if missing:
        sys.exit(f"Foto tidak ditemukan: {', '.join(missing)}")
    build()
