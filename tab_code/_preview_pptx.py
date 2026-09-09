#!/usr/bin/env python3
"""Pratinjau slide .pptx jadi PNG, tanpa PowerPoint/LibreOffice.

Menggambar ulang posisi & warna tiap shape sebagai HTML absolut lalu difoto
Chrome headless. Bukan render sempurna (pembungkusan teks bisa beda tipis),
tapi cukup untuk memeriksa tata letak: tabrakan, ruang kosong, proporsi kolom.

    python3 _preview_pptx.py tab_code_control.pptx 3 4 5
"""
import base64
import html
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
PX = 96          # piksel per inci


def inch(v):
    return Emu(v).inches


def fill_of(sh):
    try:
        f = sh.fill
        if f.type is not None and f.type == 1:
            c = f.fore_color
            return "#%02X%02X%02X" % (c.rgb[0], c.rgb[1], c.rgb[2]) if c.type == 1 else None
    except Exception:
        pass
    return None


def render(pptx, pages):
    prs = Presentation(pptx)
    sw, sh_ = inch(prs.slide_width), inch(prs.slide_height)
    out = ['<meta charset="utf-8"><style>body{margin:0;background:#3A3A3A;'
           'font-family:Arial,Helvetica,sans-serif}'
           '.s{position:relative;background:#fff;margin:14px auto;'
           f'width:{sw*PX:.0f}px;height:{sh_*PX:.0f}px;overflow:hidden}}'
           '.o{position:absolute;box-sizing:border-box}</style>']
    slides = list(prs.slides)
    for n in pages:
        s = slides[n - 1]
        out.append(f'<div class="s"><div style="position:absolute;right:6px;top:2px;'
                   f'color:#bbb;font-size:10px">slide {n}</div>')
        for shp in s.shapes:
            x, y = inch(shp.left) * PX, inch(shp.top) * PX
            w, h = inch(shp.width) * PX, inch(shp.height) * PX
            st = f"left:{x:.1f}px;top:{y:.1f}px;width:{w:.1f}px;height:{h:.1f}px;"
            if shp.has_table:                              # tabel
                rows = []
                for r in shp.table.rows:
                    tds = "".join(
                        f'<td style="border-bottom:1px solid #E3E6EC;padding:5px 8px;'
                        f'font-size:16px;vertical-align:middle">{html.escape(c.text)}</td>'
                        for c in r.cells)
                    rows.append(f"<tr>{tds}</tr>")
                out.append(f'<div class="o" style="{st}">'
                           f'<table style="width:100%;height:100%;border-collapse:collapse;table-layout:fixed">'
                           f'{"".join(rows)}</table></div>')
                continue
            if shp.shape_type == 13:                       # gambar
                b64 = base64.b64encode(shp.image.blob).decode()
                out.append(f'<img class="o" style="{st}" '
                           f'src="data:image/png;base64,{b64}">')
                continue
            bg = fill_of(shp)
            if bg:
                st += f"background:{bg};border-radius:6px;"
            body = ""
            if shp.has_text_frame and shp.text_frame.text.strip():
                for p in shp.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs)
                    if not t.strip():
                        continue
                    r0 = p.runs[0]
                    sz = (r0.font.size.pt if r0.font.size else 13) * 96 / 72   # pt -> px
                    col = "#1F2433"
                    try:
                        c = r0.font.color
                        if c and c.type == 1:
                            col = "#%02X%02X%02X" % (c.rgb[0], c.rgb[1], c.rgb[2])
                    except Exception:
                        pass
                    bold = "font-weight:700;" if r0.font.bold else ""
                    sa = (p.space_after.pt * 96/72) if p.space_after else 4
                    mono = "font-family:'Courier New',monospace;" if (
                        r0.font.name or "").startswith("Courier") else ""
                    body += (f'<div style="font-size:{sz:.1f}px;color:{col};{bold}{mono}'
                             f'line-height:1.30;margin-bottom:{sa}px;white-space:pre-wrap">'
                             f'{html.escape(t)}</div>')
                st += "padding:2px 0;"
                try:
                    if shp.text_frame.vertical_anchor == 3:
                        st += ("display:flex;flex-direction:column;"
                               "justify-content:center;")
                except Exception:
                    pass
            out.append(f'<div class="o" style="{st}">{body}</div>')
        out.append("</div>")
    return "".join(out)


if __name__ == "__main__":
    f = sys.argv[1]
    pages = [int(a) for a in sys.argv[2:]] or [1]
    tmp = ROOT / ".preview.html"
    tmp.write_text(render(ROOT / f, pages), encoding="utf-8")
    png = ROOT / ".preview.png"
    subprocess.run([CHROME, "--headless", "--disable-gpu", "--no-sandbox",
                    "--hide-scrollbars", "--allow-file-access-from-files",
                    "--virtual-time-budget=4000",
                    f"--window-size=1310,{len(pages)*(7.5*PX+14)+20:.0f}",
                    f"--screenshot={png}", f"file://{tmp}"],
                   capture_output=True, timeout=180)
    print(png)
